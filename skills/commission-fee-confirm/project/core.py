# -*- coding: utf-8 -*-
"""
佣金与两费确认文件生成器 - 核心模块（纯本地，不依赖任何 AI/LLM）

输入：社保/养老组合余额表 .xls（每组合一份）+ 可选上季应付佣金明细 .xlsx（提供席位号）
输出：每组合一个「组合名-YYYY年N季度佣金与两费确认.xlsx」，含两个 sheet：
  - 两费确认：托管费/基础管理费确认函（公式 F=D+E, H=F*20%, G=F-H）
  - 佣金明细：组合名称/券商名称/席位/应付佣金，只展示本季期末余额≠0 的券商
"""
import re
import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

# 余额表列位（本位币金额）：期初=c12, 本期借=c15, 本期贷=c18, 期末=c28
C_OPEN, C_DEBIT, C_CREDIT, C_END = 12, 15, 18, 28
TITLE_ROW = 2          # 余额表标题所在行（0-based）
DATA_START_ROW = 7     # 科目数据起始行

THIN = Side(style="thin")
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ACCT = '_ * #,##0.00_ ;_ * \\-#,##0.00_ ;_ * "-"??_ ;_ @_ '


# ---------------------------------------------------------------- 余额表解析

def parse_balance(path):
    """解析一份余额表 → dict(代码/类型/全名/期间/两费/佣金明细)"""
    df = pd.read_excel(path, sheet_name=0, header=None)
    title = str(df.iloc[TITLE_ROW, 0]).strip()
    m = re.match(r'^(HZ|YL)(\d+)(.+?)(\d{4}-\d{2}-\d{2})\s*至\s*(\d{4}-\d{2}-\d{2})', title)
    if not m:
        raise ValueError(f"无法从标题解析组合信息: {title!r}（支持 HZ/YL 前缀的余额表）")
    prefix, code, middle, start, end = m.groups()

    # 组合展示名（佣金明细的组合名称列 & 输出文件名）
    if prefix == "HZ":
        # "大成基金管理有限公司-社保基金划转1103组合" → "社保基金划转1103组合"
        display = middle.split("-", 1)[1] if "-" in middle else middle
    else:
        # "基本养老保险基金八零二组合" → "基本养老保险基金八零二"
        display = middle[:-2] if middle.endswith("组合") else middle

    q = (int(start[5:7]) - 1) // 3 + 1          # 月份→季度
    period_cn = f"{start[:4]}年{q}季度"
    period_dot = f"{start.replace('-', '.')}-{end.replace('-', '.')}"

    fee = {"mgmt": 0.0, "cust": 0.0}
    comm = {}          # 券商 → 期末余额
    for i in range(DATA_START_ROW, len(df)):
        c = df.iloc[i, 0]
        if pd.isna(c):
            continue
        s = str(c).strip()
        name, lvl = df.iloc[i, 1], df.iloc[i, 2]
        if s == "2206":
            fee["mgmt"] = float(df.iloc[i, C_CREDIT] or 0)
        elif s == "2207":
            fee["cust"] = float(df.iloc[i, C_CREDIT] or 0)
        elif s.startswith("2209.01.") and str(lvl).strip() == "3级":
            v = float(df.iloc[i, C_END] or 0)
            if abs(v) > 1e-9:                  # 只展示本季有费用的
                comm[str(name).strip()] = v

    kind = "社保划转基金" if prefix == "HZ" else "基本养老保险基金"
    return {
        "code": code, "prefix": prefix, "display": display,
        "title": f"{kind}{period_cn}两费支付确认函",
        "period_cn": period_cn, "period_dot": period_dot,
        "fee": fee, "comm": comm,
        "comm_total": sum(comm.values()),
    }


def parse_prev_commission(path):
    """解析上季佣金明细 xlsx（一个或多个 sheet）→ {组合代码(小写): [(券商,席位),...]}"""
    xl = pd.ExcelFile(path)
    seats = {}
    for sname in xl.sheet_names:
        df = pd.read_excel(path, sheet_name=sname, header=None)
        rows = []
        for i in range(len(df)):
            a = str(df.iloc[i, 0]).strip() if not pd.isna(df.iloc[i, 0]) else ""
            if a == "合计":
                break
            if i < 3 or pd.isna(df.iloc[i, 1]):
                continue
            br = str(df.iloc[i, 1]).strip()
            seat = str(df.iloc[i, 2]).strip() if not pd.isna(df.iloc[i, 2]) else ""
            rows.append((br, seat))
        seats[sname.strip().lower()] = rows
    return seats


# ---------------------------------------------------------------- sheet 渲染

def _fee_sheet(ws, info):
    for col, w in {"A": 28.5, "B": 21.6, "C": 17.9, "D": 19.1, "G": 17.9, "H": 21.9}.items():
        ws.column_dimensions[col].width = w
    for r, h in [(1, 20.25), (3, 15.0), (4, 14.25), (8, 20.25)]:
        ws.row_dimensions[r].height = h
    ws.merge_cells("A1:H1")
    ws["A1"] = info["title"]
    ws["A1"].font = Font(name="黑体", size=16)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["组合名称", "计费期间", "托管费", "基础管理费", "管理费调整情况",
               "实际支付管理费", "支付到管理人账户", "支付到风险准备金账户"]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(name="宋体", size=12, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BD

    data = {1: int(info["code"]), 2: info["period_dot"],
            3: info["fee"]["cust"], 4: info["fee"]["mgmt"],
            6: "=D5+E5", 7: "=F5-H5", 8: "=F5*0.2"}
    for j in range(1, 9):
        c = ws.cell(row=5, column=j, value=data.get(j))
        c.font = Font(name="宋体", size=11)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BD
        if j in (3, 4, 6, 7, 8):
            c.number_format = ACCT

    ws.merge_cells("A8:H8")
    ws["A8"] = "兹确认，以上表格中的数据经双方各自计算，并经双方复核后一致。"
    ws["A8"].font = Font(name="仿宋", size=16, bold=True)
    ws["A8"].alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCDEFGH":
        ws[f"{col}8"].border = Border(
            left=THIN if col == "A" else None, right=THIN if col == "H" else None,
            top=THIN, bottom=None)
    for r in (10, 12, 14):
        for col in ("B", "F"):
            c = ws[f"{col}{r}"]
            c.value = "托管行（签章）" if (col == "F" and r == 10) else \
                      ("管理人（签章）" if r == 10 else ("经办人" if r == 12 else "核对日期"))
            c.font = Font(name="宋体", size=11)
            c.alignment = Alignment(vertical="center")


def _comm_sheet(ws, info, seats):
    for col, w in {"A": 46.9, "B": 9.0, "D": 15.0, "E": 9.0}.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 14.25
    ws.row_dimensions[2].height = 15.0
    a1 = ws["A1"]; a1.value = "编制单位：大成基金管理有限公司 "
    a1.font = Font(name="宋体", size=11); a1.border = BD
    a1.alignment = Alignment(vertical="center")
    d1 = ws["D1"]; d1.value = "单位：人民币元"
    d1.font = Font(name="宋体", size=11); d1.border = BD
    d1.alignment = Alignment(vertical="center")
    ws.merge_cells("A2:D2")
    a2 = ws["A2"]; a2.value = f"期间：{info['period_cn']}"
    a2.font = Font(name="宋体", size=11, bold=True)
    a2.alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCD":
        ws[f"{col}2"].border = BD
    for j, (h, bold) in enumerate([("组合名称", True), ("券商名称", False),
                                   ("席位", False), ("应付佣金", False)], 1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = Font(name="宋体", size=11, bold=bold)
        c.alignment = Alignment(vertical="center")
        c.border = BD
    # 数据行：本季有费用的券商（保持余额表出现顺序），席位取上季明细首个
    r = 4
    for br, amt in info["comm"].items():
        seat = ""
        for pbr, pseat in seats:
            if pbr == br:
                seat = pseat
                break
        for j, v, ft, al, fmt in [
            (1, info["display"], Font(name="宋体", size=11), None, None),
            (2, br, Font(name="宋体", size=9), Alignment(horizontal="left", vertical="center"), "@"),
            (3, seat, Font(name="宋体", size=9), Alignment(horizontal="left", vertical="center"), "@"),
            (4, amt, Font(name="宋体", size=9), Alignment(horizontal="right", vertical="center"), "#,##0.00"),
        ]:
            c = ws.cell(row=r, column=j, value=v)
            c.font = ft
            if al: c.alignment = al
            if fmt: c.number_format = fmt
            c.border = BD
        r += 1
    if r > 4:  # 有数据才有合计行
        ca = ws.cell(row=r, column=1, value="合计")
        ca.font = Font(name="宋体", size=11); ca.border = BD
        for j in (2, 3):
            ws.cell(row=r, column=j).border = BD
        cd = ws.cell(row=r, column=4, value=f"=SUM(D4:D{r-1})")
        cd.font = Font(name="宋体", size=11); cd.number_format = ACCT
        cd.border = BD; cd.alignment = Alignment(vertical="center")


# ---------------------------------------------------------------- Excel 重算

def refresh_with_excel(files):
    """用本机 Excel COM 重算公式缓存值；无 Excel/COM 失败则静默跳过（文件仍可正常打开）"""
    try:
        import win32com.client
        import pythoncom
    except ImportError:
        return False
    pythoncom.CoInitialize()
    last_err = None
    for attempt in (1, 2, 3):  # 新保存文件可能被杀软短暂锁定，重试
        try:
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            try:
                for f in files:
                    f = os.path.abspath(f)
                    if not os.path.exists(f):
                        continue
                    wb = excel.Workbooks.Open(f)
                    excel.CalculateFullRebuild()
                    wb.Save()
                    wb.Close(False)
            finally:
                excel.Quit()
            return True
        except Exception as e:
            last_err = e
            import time
            time.sleep(1)
    try:
        with open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                               "refresh_error.log"), "a", encoding="utf-8") as fh:
            fh.write(f"{datetime.now()} {type(last_err).__name__}: {last_err}\n")
    except Exception:
        pass
    return False


# ---------------------------------------------------------------- 主入口

def generate(balance_files, prev_files=None, out_dir=None):
    """生成全部确认文件。返回 (结果列表, 输出文件路径列表)"""
    if out_dir is None:
        out_dir = os.path.join(os.path.expanduser("~"), "Desktop")
    os.makedirs(out_dir, exist_ok=True)

    # 上季席位结构：{组合代码(去前导零): [(券商,席位),...]}
    seat_map = {}
    for p in (prev_files or []):
        for k, v in parse_prev_commission(p).items():
            seat_map[(k.lstrip("0") or "0").lower()] = v

    results, out_files = [], []
    for bp in balance_files:
        info = parse_balance(bp)
        seats = seat_map.get((info["code"].lstrip("0") or "0").lower(), [])
        fname = f"{info['display']}-{info['period_cn']}佣金与两费确认.xlsx"
        fpath = os.path.join(out_dir, fname)

        wb = Workbook()
        ws1 = wb.active; ws1.title = "两费确认"
        _fee_sheet(ws1, info)
        ws2 = wb.create_sheet("佣金明细")
        _comm_sheet(ws2, info, seats)
        wb.save(fpath)

        results.append({
            "code": info["code"], "display": info["display"],
            "title": info["title"], "period": info["period_cn"],
            "cust_fee": info["fee"]["cust"], "mgmt_fee": info["fee"]["mgmt"],
            "comm_rows": len(info["comm"]),
            "comm_total": info["comm_total"],
            "seats_matched": sum(1 for br in info["comm"] if any(pbr == br for pbr, _ in seats)),
            "file": fname, "src": os.path.basename(bp),
        })
        out_files.append(fpath)

    refresh_with_excel(out_files)
    return results, out_files
